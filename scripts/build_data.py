"""Construye data/empresas.xlsx (fuente editable) y src/data/companies.json (datos del cliente).

Uso:
  python scripts/build_data.py            # lee el Excel si existe y genera el JSON
  python scripts/build_data.py --refresh  # reconstruye el Excel desde data/raw/*.json
                                          # conservando las columnas editables (Incluir, Rareza manual)

Los umbrales de rareza y la fórmula de ingresos deben coincidir con src/game/balance.ts.
"""
from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
XLSX = ROOT / "data" / "empresas.xlsx"
JSON_OUT = ROOT / "src" / "data" / "companies.json"

# --- Parámetros de juego (copia de src/game/balance.ts) -------------------------
THRESHOLDS_M_USD = [  # (etiqueta Excel, clave JSON, cap. mínima en millones USD, probabilidad gacha)
    ("Legendaria", "legendaria", 1_000_000, 0.01),
    ("Épica", "epica", 250_000, 0.05),
    ("Rara", "rara", 80_000, 0.12),
    ("Poco común", "poco_comun", 25_000, 0.27),
    ("Común", "comun", 0, 0.55),
]
INCOME_COEF = 0.05
INCOME_EXP = 0.75

SECTOR_ES = {
    "Technology": "Tecnología", "Financial Services": "Finanzas", "Healthcare": "Salud",
    "Consumer Cyclical": "Consumo cíclico", "Consumer Defensive": "Consumo básico",
    "Communication Services": "Comunicación", "Industrials": "Industria", "Energy": "Energía",
    "Basic Materials": "Materiales", "Real Estate": "Inmobiliario", "Utilities": "Servicios públicos",
}
COUNTRY_CODE = {
    "United States": "US", "Germany": "DE", "Spain": "ES", "United Kingdom": "GB", "France": "FR",
    "Netherlands": "NL", "Japan": "JP", "Switzerland": "CH", "China": "CN", "Ireland": "IE",
    "Sweden": "SE", "Italy": "IT", "India": "IN", "Australia": "AU", "Denmark": "DK", "Canada": "CA",
    "South Korea": "KR", "Finland": "FI", "Hong Kong": "HK", "Saudi Arabia": "SA", "Belgium": "BE",
    "Norway": "NO", "Uruguay": "UY", "Luxembourg": "LU", "Taiwan": "TW", "Austria": "AT",
    "Portugal": "PT", "Singapore": "SG", "Brazil": "BR", "Mexico": "MX", "Israel": "IL",
    "Bermuda": "BM", "New Zealand": "NZ", "Macau": "MO", "South Africa": "ZA", "Greece": "GR", "Cayman Islands": "KY", "Jersey": "JE",
}
LABEL_TO_KEY = {label: key for label, key, _, _ in THRESHOLDS_M_USD}

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BOLD = Font(name=FONT, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")

HEADERS = [
    "Ticker", "Nombre", "País", "Código país", "Sector", "Industria", "Bolsa", "Moneda",
    "Cap. local (M)", "Tipo cambio a USD", "Cap. USD (M)", "Rareza", "Ingresos local (M)",
    "Beneficio local (M)", "Ingresos USD (M)", "Beneficio USD (M)", "Rent. dividendo", "PER",
    "Índices", "Incluir", "Rareza manual", "Ingresos/s en juego",
]
COL = {h: get_column_letter(i + 1) for i, h in enumerate(HEADERS)}


def rarity_key(cap_m_usd: float) -> str:
    for _, key, minimum, _ in THRESHOLDS_M_USD:
        if cap_m_usd >= minimum:
            return key
    return "comun"


def read_overrides() -> dict[str, tuple[str, str]]:
    """Devuelve {ticker: (Incluir, Rareza manual)} del Excel actual, si existe."""
    if not XLSX.exists():
        return {}
    wb = load_workbook(XLSX, read_only=True)
    ws = wb["Empresas"]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            out[row[0]] = (row[HEADERS.index("Incluir")] or "SI", row[HEADERS.index("Rareza manual")] or "")
    wb.close()
    return out


def build_xlsx() -> None:
    rows = json.loads((RAW / "yahoo_info.json").read_text(encoding="utf-8"))
    fx = json.loads((RAW / "fx.json").read_text(encoding="utf-8"))
    overrides = read_overrides()
    rows = [r for r in rows if "error" not in r and r.get("marketCap")]
    rows.sort(key=lambda r: -(r["marketCap"] * fx.get(norm_cur(r.get("currency")), 0)))

    wb = Workbook()

    # --- Leyenda ---------------------------------------------------------------
    ws = wb.active
    ws.title = "Leyenda"
    lines = [
        ("CapitalDex — Universo de empresas", BOLD),
        (f"Generado el {date.today().isoformat()} desde Yahoo Finance (yfinance). Datos con retraso; solo para diseño del juego.", None),
        ("", None),
        ("Celdas editables (texto azul, fondo amarillo):", BOLD),
        ("  Empresas!Incluir — SI/NO. NO saca la empresa del juego.", None),
        ("  Empresas!Rareza manual — fuerza una rareza (Legendaria, Épica, Rara, Poco común, Común). Vacío = automática.", None),
        ("  Umbrales!B2:B6 — capitalización mínima (M USD) de cada rareza.", None),
        ("  Umbrales!C2:C6 — probabilidad de cada rareza en el gacha (deben sumar 100 %).", None),
        ("  Umbrales!B9:B10 — coeficiente y exponente de la fórmula de ingresos.", None),
        ("  FX!B — tipo de cambio a USD.", None),
        ("", None),
        ("Después de editar, ejecuta:  npm run data   (o python scripts/build_data.py) para regenerar el JSON del juego.", None),
        ("Los umbrales y la fórmula de ingresos también están en src/game/balance.ts; mantenlos iguales.", None),
        ("Nota: la moneda GBp (peniques) se normaliza a GBP porque Yahoo da la capitalización ya en libras.", None),
    ]
    for i, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font or Font(name=FONT)
    ws.column_dimensions["A"].width = 120

    # --- FX --------------------------------------------------------------------
    wfx = wb.create_sheet("FX")
    wfx.append(["Moneda", "USD por unidad", "Fuente"])
    for cur, rate in sorted(fx.items()):
        wfx.append([cur, rate, f"Yahoo Finance {cur}USD=X, cierre {date.today().isoformat()}"])
    style_header(wfx, 3)
    for r in range(2, wfx.max_row + 1):
        wfx[f"B{r}"].font = BLUE
        wfx[f"B{r}"].fill = INPUT_FILL
        wfx[f"B{r}"].number_format = "0.000000"
    wfx.column_dimensions["A"].width = 10
    wfx.column_dimensions["B"].width = 16
    wfx.column_dimensions["C"].width = 44

    # --- Umbrales ----------------------------------------------------------------
    wu = wb.create_sheet("Umbrales")
    wu.append(["Rareza", "Cap. mínima (M USD)", "Probabilidad gacha", "Empresas", "% del total"])
    for i, (label, _, minimum, prob) in enumerate(THRESHOLDS_M_USD, start=2):
        wu.append([label, minimum, prob, None, None])
        wu[f"B{i}"].font, wu[f"B{i}"].fill, wu[f"B{i}"].number_format = BLUE, INPUT_FILL, "#,##0"
        wu[f"C{i}"].font, wu[f"C{i}"].fill, wu[f"C{i}"].number_format = BLUE, INPUT_FILL, "0.0%"
        wu[f"D{i}"] = f'=COUNTIFS(Empresas!${COL["Rareza"]}:${COL["Rareza"]},A{i},Empresas!${COL["Incluir"]}:${COL["Incluir"]},"SI")'
        wu[f"E{i}"] = f"=IF(SUM($D$2:$D$6)=0,0,D{i}/SUM($D$2:$D$6))"
        wu[f"E{i}"].number_format = "0.0%"
    wu["A7"], wu["C7"], wu["D7"] = "Total", "=SUM(C2:C6)", "=SUM(D2:D6)"
    wu["C7"].number_format = "0.0%"
    wu["A7"].font = BOLD
    wu["A9"], wu["B9"] = "Coeficiente ingresos", INCOME_COEF
    wu["A10"], wu["B10"] = "Exponente ingresos", INCOME_EXP
    wu["A11"] = "ingresos/s = coeficiente × (Cap. USD en miles de millones) ^ exponente"
    for ref in ("B9", "B10"):
        wu[ref].font, wu[ref].fill = BLUE, INPUT_FILL
    wu["B2"].comment = Comment("Umbrales del GDD (sección 3), ajustados el 15-oct-2026 al reparto real del universo. Objetivo aproximado: 3/6/20/40/30 %.", "CapitalDex")
    style_header(wu, 5)
    for col, w in zip("ABCDE", (22, 22, 20, 12, 12)):
        wu.column_dimensions[col].width = w

    # --- Empresas ----------------------------------------------------------------
    we = wb.create_sheet("Empresas", 1)
    we.append(HEADERS)
    for i, r in enumerate(rows, start=2):
        cur = norm_cur(r.get("currency"))
        incl, manual = overrides.get(r["ticker"], ("SI", ""))
        country = r.get("country") or ""
        we.append([
            r["ticker"], r.get("longName") or r.get("shortName") or r["ticker"], country,
            COUNTRY_CODE.get(country, "UN"), SECTOR_ES.get(r.get("sector") or "", r.get("sector") or "") or "Otros",
            r.get("industry") or "", r.get("exchange") or "", cur,
            r["marketCap"] / 1e6,
            f'=INDEX(FX!$B:$B,MATCH({COL["Moneda"]}{i},FX!$A:$A,0))',
            f'={COL["Cap. local (M)"]}{i}*{COL["Tipo cambio a USD"]}{i}',
            rarity_formula(i),
            (r.get("totalRevenue") or 0) / 1e6 or None,
            (r.get("netIncomeToCommon") or 0) / 1e6 or None,
            f'=IF({COL["Ingresos local (M)"]}{i}="","",{COL["Ingresos local (M)"]}{i}*{COL["Tipo cambio a USD"]}{i})',
            f'=IF({COL["Beneficio local (M)"]}{i}="","",{COL["Beneficio local (M)"]}{i}*{COL["Tipo cambio a USD"]}{i})',
            (r["dividendYield"] / 100) if r.get("dividendYield") is not None else None,  # Yahoo lo da en %
            r.get("trailingPE"),
            ", ".join(r.get("indices") or []), incl, manual,
            f'=Umbrales!$B$9*({COL["Cap. USD (M)"]}{i}/1000)^Umbrales!$B$10',
        ])
        for h in ("Incluir", "Rareza manual"):
            we[f"{COL[h]}{i}"].font, we[f"{COL[h]}{i}"].fill = BLUE, INPUT_FILL
        for h in ("Cap. local (M)", "Cap. USD (M)", "Ingresos local (M)", "Beneficio local (M)", "Ingresos USD (M)", "Beneficio USD (M)"):
            we[f"{COL[h]}{i}"].number_format = "#,##0"
        we[f"{COL['Tipo cambio a USD']}{i}"].number_format = "0.0000"
        we[f"{COL['Rent. dividendo']}{i}"].number_format = "0.00%"
        we[f"{COL['PER']}{i}"].number_format = "0.0"
        we[f"{COL['Ingresos/s en juego']}{i}"].number_format = "0.00"
    style_header(we, len(HEADERS))
    we.freeze_panes = "C2"
    we.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{we.max_row}"
    widths = {"Ticker": 12, "Nombre": 34, "País": 16, "Sector": 18, "Industria": 30, "Bolsa": 8,
              "Índices": 26, "Rareza manual": 14, "Ingresos/s en juego": 12}
    for h in HEADERS:
        we.column_dimensions[COL[h]].width = widths.get(h, 13)
    we[f"{COL['Rent. dividendo']}1"].comment = Comment("Fracción (0,02 = 2 %). Fuente: Yahoo Finance, campo dividendYield (que viene en %, dividido entre 100).", "CapitalDex")
    we[f"{COL['Cap. local (M)']}1"].comment = Comment("Millones en moneda local. Fuente: Yahoo Finance, campo marketCap.", "CapitalDex")

    wb.save(XLSX)
    print(f"Excel escrito: {XLSX.relative_to(ROOT)} ({len(rows)} empresas)")


def rarity_formula(i: int) -> str:
    k, m = f"{COL['Cap. USD (M)']}{i}", f"{COL['Rareza manual']}{i}"
    return (f'=IF({m}<>"",{m},IF({k}>=Umbrales!$B$2,"Legendaria",IF({k}>=Umbrales!$B$3,"Épica",'
            f'IF({k}>=Umbrales!$B$4,"Rara",IF({k}>=Umbrales!$B$5,"Poco común","Común")))))')


def norm_cur(cur: str | None) -> str:
    return "GBP" if cur == "GBp" else (cur or "USD")


def style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.font.name != FONT or cell.font.color is None:
                cell.font = Font(name=FONT, bold=cell.font.bold, color=cell.font.color)


def build_json() -> None:
    wb = load_workbook(XLSX, read_only=True)
    fx = {r[0]: r[1] for r in wb["FX"].iter_rows(min_row=2, values_only=True) if r[0]}
    ws = wb["Empresas"]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(HEADERS, row))
        if not rec["Ticker"] or (rec["Incluir"] or "SI").upper() != "SI":
            continue
        rate = fx.get(rec["Moneda"], None)
        if not rate:
            print(f"  aviso: sin tipo de cambio para {rec['Ticker']} ({rec['Moneda']}), se omite")
            continue
        cap_m = float(rec["Cap. local (M)"]) * rate
        item = {
            "id": rec["Ticker"], "name": rec["Nombre"] or rec["Ticker"], "sector": rec["Sector"] or "Otros",
            "industry": rec["Industria"] or "", "country": rec["País"] or "", "countryCode": rec["Código país"] or "UN",
            "exchange": rec["Bolsa"] or "", "currency": rec["Moneda"], "marketCapUsd": round(cap_m * 1e6),
            "revenueUsd": round(float(rec["Ingresos local (M)"]) * rate * 1e6) if rec["Ingresos local (M)"] else None,
            "netIncomeUsd": round(float(rec["Beneficio local (M)"]) * rate * 1e6) if rec["Beneficio local (M)"] else None,
            "dividendYield": rec["Rent. dividendo"], "pe": rec["PER"],
            "indices": [s.strip() for s in (rec["Índices"] or "").split(",") if s.strip()],
        }
        manual = (rec["Rareza manual"] or "").strip()
        if manual:
            if manual not in LABEL_TO_KEY:
                sys.exit(f"Rareza manual desconocida en {rec['Ticker']}: {manual!r}")
            item["rarityOverride"] = LABEL_TO_KEY[manual]
        out.append(item)
    wb.close()
    out.sort(key=lambda c: -c["marketCapUsd"])
    JSON_OUT.parent.mkdir(parents=True, exist_ok=True)
    JSON_OUT.write_text(json.dumps(out, ensure_ascii=False, indent=0), encoding="utf-8")

    counts: dict[str, int] = {}
    for c in out:
        k = c.get("rarityOverride") or rarity_key(c["marketCapUsd"] / 1e6)
        counts[k] = counts.get(k, 0) + 1
    print(f"JSON escrito: {JSON_OUT.relative_to(ROOT)} ({len(out)} empresas)")
    for _, key, _, _ in THRESHOLDS_M_USD:
        n = counts.get(key, 0)
        print(f"  {key:11s} {n:4d}  {n / len(out):5.1%}")


if __name__ == "__main__":
    if "--refresh" in sys.argv or not XLSX.exists():
        build_xlsx()
    build_json()
